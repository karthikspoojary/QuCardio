import docx
from docx.shared import Inches
import os

doc = docx.Document()
doc.add_heading('QuCardio System Design Images', 0)

images = [
    ("System Architecture", "fig4_1_system_architecture.png"),
    ("Use Case Diagram", "fig4_2_use_case.png"),
    ("Data Flow Diagram", "fig4_3_dfd.png"),
    ("Class Diagram", "fig4_4_class_diagram.png"),
    ("Sequence Diagram", "fig4_5_sequence.png")
]

base_dir = "/home/karthik/projects/QuCardio/Doc/report/extracted_images/"

for title, img_file in images:
    doc.add_heading(title, level=1)
    img_path = os.path.join(base_dir, img_file)
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6.0))
    else:
        doc.add_paragraph(f"Image not found: {img_file}")

doc.save("/home/karthik/projects/QuCardio/Doc/report/qucardio images.docx")
print("Docx created successfully.")
