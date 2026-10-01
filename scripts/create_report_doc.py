import docx
from docx.shared import Inches
import os

def add_image_to_doc(doc, img_path, caption=""):
    if os.path.exists(img_path):
        doc.add_paragraph(caption)
        doc.add_picture(img_path, width=Inches(6.0))
        print(f"Added {img_path}")
    else:
        print(f"MISSING: {img_path}")

report_doc = docx.Document()
report_doc.add_heading('QuCardio Report Figures (v9)', 0)

report_images = [
    ("Figure 4.1: System architecture showing the offline training and online inference paths", "Doc/report/extracted_images/qucardio_architecture_ieee (2).png"),
    ("Figure 4.2: Use case diagram", "Doc/report/extracted_images/qucardio_use_case_v2_bw (1).png"),
    ("Figure 4.3: Level-1 data flow diagram", "Doc/report/extracted_images/fig4_3_dfd_ieee (1).png"),
    ("Figure 4.4: Class diagram", "Doc/report/extracted_images/fig4_4_class_diagram_ieee (1).png"),
    ("Figure 4.5: Sequence diagram for the POST /predict request", "Doc/report/extracted_images/fig4_5_sequence_diagram.png"),
    ("Figure 5.1: Raw clinical ECG scan and the resulting 340x340 preprocessed output", "Doc/screenshots/5.1_preprocessing.png"),
    ("Figure 5.2: ZZFeatureMap circuit with nine qubits, reps = 2 and circular entanglement", "results/paper/main/zzfeaturemap_circuit.png"),
    ("Figure 5.3: Main upload interface with the drag-and-drop area and the three model selector buttons", "Doc/screenshots/doc1_image1.png"),
    ("Figure 5.4: Diagnosis result screen showing the predicted class, the confidence bar, the per-class probability bars, the preprocessed ECG thumbnail and the clinical recommendation", "Doc/screenshots/doc1_image2.png"),
    ("Figure 5.5: Model performance tab with the accuracy comparison chart for the three classifiers", "Doc/screenshots/doc1_image4.png"),
    ("Figure 7.1: Accuracy comparison across all models", "results/exploratory/comparison_all_models.png"),
    ("Figure 7.2: QSVC confusion matrix", "results/paper/main/confusion_matrix_qsvc_tuned.png"),
    ("Figure 7.3: Pegasos QSVC confusion matrix", "results/paper/main/confusion_matrix_pegasos_tuned.png"),
    ("Figure 7.4: Feature encoding range ablation", "results/paper/ablation/ablation_encoding.png"),
    ("Figure 7.5: Entanglement topology ablation", "results/paper/ablation/ablation_entanglement.png"),
    ("Figure 7.6: Pauli feature map ablation", "results/paper/ablation/ablation_feature_maps.png"),
    ("Figure 7.7: ROC curves for all models and classes", "results/paper/main/roc_curves_all_models.png"),
    ("Figure 7.8: Confidence calibration curves", "results/paper/main/calibration_all_models.png"),
    ("Figure 7.9: Augmentation robustness study", "results/paper/ablation/augmentation_robustness.png"),
    ("Figure 7.10: QSVC cross-dataset confusion matrix", "results/paper/cross_dataset/jain_dataset2/cm_qsvc.png"),
    ("Figure 7.11: Classical SVM cross-dataset confusion matrix", "results/paper/cross_dataset/jain_dataset2/cm_svm.png")
]

for caption, path in report_images:
    add_image_to_doc(report_doc, path, caption)

report_doc.save('Doc/Report/doc1.docx')
print("Done creating updated report doc!")
